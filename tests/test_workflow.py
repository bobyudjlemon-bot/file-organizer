import json
import os
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook, Workbook
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from file_organizer.core import preview, apply, extract, digest


@pytest.fixture
def source(tmp_path):
    folder = tmp_path / '原件'
    folder.mkdir()
    return folder


def test_full_workflow_preserves_originals(source, tmp_path):
    (source / '合同.txt').write_text('合同金额100元', encoding='utf-8')
    (source / '副本.txt').write_text('合同金额100元', encoding='utf-8')
    (source / '照片.JPG').write_bytes(b'image bytes')
    before = {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in source.iterdir()}
    rules = {'group_by': ['category', 'year'], 'rename': '{category}_{index}', 'rules': [{'category': '合同', 'content_contains': ['合同']}]}
    plan = preview(source, tmp_path / 'preview', rules)
    assert len(plan['files']) == 3
    assert len(plan['duplicates']) == 1
    assert any(r['warning'] for r in plan['files'])
    assert not (tmp_path / 'preview' / 'files').exists()
    result = apply(tmp_path / 'preview/plan.json', tmp_path / 'result', True)
    assert result['copied'] == 3 and result['failed'] == 0
    for row in plan['files']:
        target = tmp_path / 'result/files' / row['target']
        assert digest(target) == row['sha256']
        assert target.suffix == Path(row['source']).suffix
    assert before == {p.name: (p.read_bytes(), p.stat().st_mtime_ns) for p in source.iterdir()}
    book = load_workbook(tmp_path / 'result/inventory.xlsx')
    assert book['文件清单'].max_row == 4
    assert book['重复文件'].max_row == 3
    with zipfile.ZipFile(tmp_path / 'result/organized.zip') as z:
        assert z.testzip() is None
        assert len(z.namelist()) == 6


def test_name_collisions(source, tmp_path):
    for directory in ('a', 'b'):
        (source / directory).mkdir()
        (source / directory / 'Same.txt').write_text(directory)
    plan = preview(source, tmp_path / 'p', {'rename': '统一名字'})
    assert len({r['target'].casefold() for r in plan['files']}) == 2
    assert all(r['target'].endswith('.txt') for r in plan['files'])


def test_changed_original_fails_without_bad_copy(source, tmp_path):
    path = source / 'a.txt'
    path.write_text('before')
    preview(source, tmp_path / 'p', {})
    path.write_text('after')
    result = apply(tmp_path / 'p/plan.json', tmp_path / 'r', True)
    assert result['failed'] == 1 and result['copied'] == 0
    assert path.read_text() == 'after'
    assert not list((tmp_path / 'r/files').rglob('*.txt'))
    assert '预览不一致' in (tmp_path / 'r/errors.json').read_text()


def test_confirmation_and_existing_output(source, tmp_path):
    (source / 'a.txt').write_text('a')
    preview(source, tmp_path / 'p', {})
    with pytest.raises(ValueError, match='确认'):
        apply(tmp_path / 'p/plan.json', tmp_path / 'r')
    with pytest.raises(FileExistsError):
        preview(source, tmp_path / 'p', {})
    apply(tmp_path / 'p/plan.json', tmp_path / 'r', True)
    with pytest.raises(FileExistsError):
        apply(tmp_path / 'p/plan.json', tmp_path / 'r', True)


@pytest.mark.parametrize('target', ['../escape.txt', '/tmp/escape.txt', 'folder/../../escape.txt', 'a\\b'])
def test_malicious_plan_rejected(source, tmp_path, target):
    (source / 'a.txt').write_text('a')
    plan = preview(source, tmp_path / 'p', {})
    plan['files'][0]['target'] = target
    path = tmp_path / 'p/plan.json'
    path.write_text(json.dumps(plan))
    with pytest.raises(ValueError):
        apply(path, tmp_path / 'r', True)
    assert not (tmp_path / 'r').exists()


def test_links_and_special_files_skipped(source, tmp_path):
    outside = tmp_path / 'outside.txt'
    outside.write_text('secret')
    (source / 'link').symlink_to(outside)
    (source / 'dirlink').symlink_to(tmp_path, target_is_directory=True)
    os.mkfifo(source / 'pipe')
    plan = preview(source, tmp_path / 'p', {})
    assert len(plan['errors']) == 3 and not plan['files']


def test_overlap_rejected(source):
    with pytest.raises(ValueError):
        preview(source, source / 'output', {})


def test_excel_formula_is_literal(source, tmp_path):
    (source / '=1+1.txt').write_text('hello')
    preview(source, tmp_path / 'p', {})
    book = load_workbook(tmp_path / 'p/preview.xlsx')
    assert book['文件清单']['A2'].data_type == 's'


def test_extraction_office(source):
    doc = Document()
    doc.add_paragraph('采购合同')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = '表格内容'
    doc.save(source / 'a.docx')
    assert '表格内容' in extract(source / 'a.docx')
    book = Workbook()
    book.active.append(['采购合同', 100])
    book.save(source / 'a.xlsx')
    assert '采购合同' in extract(source / 'a.xlsx')
    ppt = Presentation()
    slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = '采购合同'
    ppt.save(source / 'a.pptx')
    assert '采购合同' in extract(source / 'a.pptx')


def test_pdf_empty_and_encrypted_logged(source, tmp_path):
    pdf = PdfWriter()
    pdf.add_blank_page(width=100, height=100)
    pdf.write(source / 'blank.pdf')
    pdf.encrypt('password')
    pdf.write(source / 'locked.pdf')
    plan = preview(source, tmp_path / 'p', {'rules': [{'category': '合同', 'content_contains': ['合同']}]})
    assert len(plan['files']) == 2
    assert all(row['warning'] for row in plan['files'])
    assert any('加密' in row['warning'] for row in plan['files'])


def test_filename_date_rules_and_priority(source, tmp_path):
    p = source / '账单.TXT'
    p.write_text('内容')
    os.utime(p, (1704067200, 1704067200))
    plan = preview(source, tmp_path / 'p', {'group_by': ['year', 'month', 'category'], 'rules': [{'category': '账单', 'filename_contains': ['账单'], 'extensions': ['.txt'], 'date_from': '2024-01-01', 'date_to': '2024-12-31'}, {'category': '后续'}]})
    assert plan['files'][0]['target'] == '2024/2024-01/账单/账单.TXT'


def test_removed_file_recorded(source, tmp_path):
    p = source / 'a.txt'
    p.write_text('a')
    preview(source, tmp_path / 'p', {})
    p.unlink()
    result = apply(tmp_path / 'p/plan.json', tmp_path / 'r', True)
    assert result['failed'] == 1


def test_bad_content_falls_back(source, tmp_path):
    (source / 'broken.docx').write_bytes(b'not a zip')
    plan = preview(source, tmp_path / 'p', {'rules': [{'category': '合同', 'content_contains': ['合同']}]})
    assert plan['files'][0]['category'] == 'Word'
    assert plan['files'][0]['warning']
    assert apply(tmp_path / 'p/plan.json', tmp_path / 'r', True)['copied'] == 1


@pytest.mark.parametrize('rules', [{'rename': '{stem.__class__}'}, {'group_by': ['unknown']}, {'rules': [{'category': 'x', 'content_contains': ''}]}])
def test_bad_rules(source, tmp_path, rules):
    with pytest.raises(ValueError):
        preview(source, tmp_path / 'p', rules)
    assert not (tmp_path / 'p').exists()


def test_import_zip(tmp_path):
    from file_organizer.core import import_zip
    archive = tmp_path / 'input.zip'
    with zipfile.ZipFile(archive, 'w') as z:
        z.writestr('目录/合同.txt', '合同')
    result = import_zip(archive, tmp_path / 'imported')
    assert result['imported'] == 1 and result['failed'] == 0
    assert (tmp_path / 'imported/目录/合同.txt').read_text() == '合同'
    with pytest.raises(FileExistsError):
        import_zip(archive, tmp_path / 'imported')


@pytest.mark.parametrize('filename', ['../outside', '/tmp/outside', 'C:/outside', 'a\\b'])
def test_zip_traversal(tmp_path, filename):
    from file_organizer.core import import_zip
    archive = tmp_path / 'input.zip'
    with zipfile.ZipFile(archive, 'w') as z:
        z.writestr(filename, 'bad')
    with pytest.raises(ValueError):
        import_zip(archive, tmp_path / 'imported')
    assert not (tmp_path / 'imported').exists()


def test_zip_link(tmp_path):
    from file_organizer.core import import_zip
    archive = tmp_path / 'input.zip'
    info = zipfile.ZipInfo('link')
    info.create_system = 3
    info.external_attr = 0o120777 << 16
    with zipfile.ZipFile(archive, 'w') as z:
        z.writestr(info, '/etc/passwd')
    with pytest.raises(ValueError):
        import_zip(archive, tmp_path / 'imported')


def test_pdf_text(source):
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 12 Tf 20 100 Td (Invoice contract) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    writer.write(source / 'text.pdf')
    assert 'Invoice contract' in extract(source / 'text.pdf')


def test_cli_end_to_end(tmp_path):
    import subprocess
    import sys
    archive = tmp_path / 'upload.zip'
    with zipfile.ZipFile(archive, 'w') as z:
        z.writestr('客户/合同.txt', '合同内容')
        z.writestr('客户/合同副本.txt', '合同内容')
    def run(*args):
        return subprocess.run([sys.executable, '-m', 'file_organizer', *map(str, args)], capture_output=True, text=True)
    assert run('import-zip', archive, '--output', tmp_path / 'input').returncode == 0
    result = run('preview', tmp_path / 'input', '--output', tmp_path / 'preview')
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)['duplicate_groups'] == 1
    assert run('apply', tmp_path / 'preview/plan.json', '--output', tmp_path / 'result').returncode == 1
    assert not (tmp_path / 'result').exists()
    result = run('apply', tmp_path / 'preview/plan.json', '--output', tmp_path / 'result', '--confirm')
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)['copied'] == 2
    with zipfile.ZipFile(tmp_path / 'result/organized.zip') as z:
        assert z.testzip() is None
        assert len([name for name in z.namelist() if name.startswith('files/')]) == 2
