"""Preview first; copy only into a newly created output directory."""
import argparse
import hashlib
import json
import os
import re
import stat
import sys
import zipfile
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook

TYPES = {
    'Word': {'.doc', '.docx', '.odt', '.rtf'},
    'PDF': {'.pdf'},
    'Excel': {'.xls', '.xlsx', '.xlsm', '.csv'},
    'PPT': {'.ppt', '.pptx'},
    '图片': {'.jpg', '.jpeg', '.png', '.gif', '.bmp', '.tif', '.tiff', '.webp', '.heic', '.svg'},
    '压缩包': {'.zip', '.rar', '.7z', '.tar', '.gz', '.bz2', '.xz'},
    '文本': {'.txt', '.md', '.json', '.log'},
}
LIMIT = 200_000


def regular_open(path):
    fd = os.open(path, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK)
    if not stat.S_ISREG(os.fstat(fd).st_mode):
        os.close(fd)
        raise ValueError('不是普通文件')
    return os.fdopen(fd, 'rb')


def digest(path):
    h = hashlib.sha256()
    with regular_open(path) as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def component(value):
    value = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', str(value)).strip(' .')
    value = value[:80]
    if not value or value in {'.', '..'}:
        raise ValueError('名称不能为空')
    if value.split('.')[0].upper() in {'CON', 'PRN', 'AUX', 'NUL', *(f'COM{i}' for i in range(1, 10)), *(f'LPT{i}' for i in range(1, 10))}:
        value = '_' + value
    return value


def extract(path, ocr=False):
    """Bounded text extraction. Never execute macros or unpack user archives."""
    ext = path.suffix.lower()
    if ocr:
        from .ocr import IMAGE_EXTENSIONS, extract_ocr
        if ext in IMAGE_EXTENSIONS or ext == '.pdf':
            text = extract_ocr(path)
            if len(text) > LIMIT:
                raise ValueError('OCR文本超过20万字符')
            return text
    if path.stat().st_size > 30 * 1024 * 1024:
        raise ValueError('超过30MB，跳过内容提取；仍按名称/类型/日期整理')
    if ext in {'.docx', '.xlsx', '.xlsm', '.pptx'}:
        with zipfile.ZipFile(path) as archive:
            if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Office解压大小超过100MB，跳过内容提取')
    if ext in {'.txt', '.md', '.csv', '.json', '.log'}:
        raw = path.read_bytes()
        try:
            text = raw.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = raw.decode('gb18030')
    elif ext == '.docx':
        from docx import Document
        doc = Document(path)
        text = '\n'.join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('PDF已加密，未读取内容')
        if len(reader.pages) > 100:
            raise ValueError('PDF超过100页，跳过内容提取')
        text = '\n'.join(p.extract_text() or '' for p in reader.pages)
        if not text.strip():
            raise ValueError('PDF无可提取文本，可能需要OCR')
    elif ext in {'.xlsx', '.xlsm', '.xls'}:
        chunks = []
        length = 0
        if ext == '.xls':
            import xlrd
            book = xlrd.open_workbook(path, on_demand=True)
            try:
                for sheet in book.sheets():
                    for row in range(sheet.nrows):
                        chunks.append(' '.join(str(v) for v in sheet.row_values(row)))
                        length += len(chunks[-1])
                        if length > LIMIT:
                            raise ValueError('表格文本超过20万字符，跳过内容提取')
            finally:
                book.release_resources()
        else:
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            try:
                for sheet in book:
                    for row in sheet.iter_rows(values_only=True):
                        chunks.append(' '.join(str(v) for v in row if v is not None))
                        length += len(chunks[-1])
                        if length > LIMIT:
                            raise ValueError('表格文本超过20万字符，跳过内容提取')
            finally:
                book.close()
        text = '\n'.join(chunks)
    elif ext == '.pptx':
        from pptx import Presentation
        text = '\n'.join(shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame)
    else:
        raise ValueError('该格式不支持内容提取（图片/OCR、旧版Word/PPT及压缩包等）；仍支持其他分类')
    if len(text) > LIMIT:
        raise ValueError('文本超过20万字符，跳过内容提取')
    return text


def validate_rules(rules):
    if set(rules) - {'group_by', 'rules', 'rename', 'timezone', 'ocr'}:
        raise ValueError('未知配置字段')
    if 'ocr' in rules and not isinstance(rules['ocr'], bool):
        raise ValueError('ocr必须是true或false')
    groups = rules.get('group_by', ['category'])
    if not isinstance(groups, list) or any(x not in {'category', 'type', 'year', 'month'} for x in groups):
        raise ValueError('group_by仅允许category/type/year/month')
    ZoneInfo(rules.get('timezone', 'Asia/Shanghai'))
    for rule in rules.get('rules', []):
        if set(rule) - {'category', 'filename_contains', 'content_contains', 'extensions', 'date_from', 'date_to'}:
            raise ValueError('未知分类规则字段')
        component(rule['category'])
        for key in ('filename_contains', 'content_contains', 'extensions'):
            if key in rule and (not isinstance(rule[key], list) or not rule[key] or not all(isinstance(x, str) and x for x in rule[key])):
                raise ValueError(f'{key}必须是非空字符串列表')
        for key in ('date_from', 'date_to'):
            if key in rule:
                datetime.strptime(rule[key], '%Y-%m-%d')
    rename = rules.get('rename')
    if rename is not None:
        if not isinstance(rename, str):
            raise ValueError('rename必须是字符串模板')
        # Only simple documented placeholders; no attribute/index access.
        if re.sub(r'\{(stem|index|date|category|type)\}', '', rename).find('{') >= 0 or '}' in re.sub(r'\{(stem|index|date|category|type)\}', '', rename):
            raise ValueError('不支持的重命名占位符')


def matches(rule, row, text):
    for key, haystack in [('filename_contains', row['name']), ('content_contains', text)]:
        if key in rule and not any(word.casefold() in haystack.casefold() for word in rule[key]):
            return False
    if 'extensions' in rule and row['extension'] not in [x.lower() for x in rule['extensions']]:
        return False
    day = row['date'][:10]
    return not (day < rule.get('date_from', day) or day > rule.get('date_to', day))


def save_json(path, obj):
    with path.open('x', encoding='utf-8') as stream:
        json.dump(obj, stream, ensure_ascii=False, indent=2)


def report(path, rows, duplicates, errors):
    book = Workbook()
    sheet = book.active
    sheet.title = '文件清单'
    columns = [('name', '文件名'), ('source', '原路径'), ('extension', '格式'), ('size', '大小(字节)'), ('date', '修改日期'), ('category', '分类结果'), ('target', '整理后相对路径'), ('sha256', 'SHA256'), ('status', '状态'), ('warning', '原因/内容提取提示')]
    sheet.append([title for _, title in columns])
    for row in rows:
        sheet.append([row.get(key, '') for key, _ in columns])
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    dup = book.create_sheet('重复文件')
    dup.append(['重复组SHA256', '路径'])
    for key, paths in duplicates.items():
        for name in paths:
            dup.append([key, name])
    err = book.create_sheet('失败记录')
    err.append(['路径', '原因'])
    for item in errors:
        err.append([item['path'], item['reason']])
    for tab in book:
        for cells in tab:
            for cell in cells:
                if isinstance(cell.value, str):
                    # Literal strings prevent Excel formula injection via filenames.
                    cell.value = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', cell.value)
                    cell.data_type = 's'
        for col in tab.columns:
            tab.column_dimensions[col[0].column_letter].width = 28
    with path.open('xb') as stream:
        book.save(stream)


def separate(source, output):
    if output == source or source in output.parents or output in source.parents:
        raise ValueError('输出目录与原始目录必须分开，不能互相包含')


def preview(source, output, rules):
    source, output = Path(source).resolve(), Path(output).resolve()
    if not source.is_dir():
        raise ValueError('输入目录不存在')
    separate(source, output)
    validate_rules(rules)
    output.mkdir(parents=True, exist_ok=False)
    rows, errors, used = [], [], set()
    hashes = defaultdict(list)
    def walk_error(exc):
        errors.append({'path': str(exc.filename), 'reason': str(exc)})
    for root, dirs, files in os.walk(source, followlinks=False, onerror=walk_error):
        dirs.sort()
        for name in dirs[:]:
            path = Path(root) / name
            if path.is_symlink():
                dirs.remove(name)
                errors.append({'path': str(path), 'reason': '跳过符号链接目录'})
        for name in sorted(files):
            path = Path(root) / name
            try:
                if path.is_symlink():
                    raise ValueError('跳过符号链接')
                info = path.stat()
                checksum = digest(path)
                ext = path.suffix.lower()
                kind = next((k for k, v in TYPES.items() if ext in v), '其他')
                day = datetime.fromtimestamp(info.st_mtime, ZoneInfo(rules.get('timezone', 'Asia/Shanghai'))).isoformat(timespec='seconds')
                row = dict(name=name, source=str(path), extension=ext, size=info.st_size, date=day, category=kind, type=kind, sha256=checksum, status='预览', warning='')
                text = ''
                if any('content_contains' in rule for rule in rules.get('rules', [])):
                    try:
                        text = extract(path, ocr=rules.get('ocr', False))
                        if rules.get('ocr', False):
                            row['warning'] = '已启用OCR辅助内容分类；机器识别可能出错，请核对预览'
                    except Exception as exc:
                        row['warning'] = str(exc)
                for rule in rules.get('rules', []):
                    if matches(rule, row, text):
                        row['category'] = component(rule['category'])
                        break
                values = dict(stem=path.stem, index=f'{len(rows)+1:04d}', date=day[:10], category=row['category'], type=kind)
                stem = component(rules['rename'].format(**values)) if rules.get('rename') else component(path.stem)
                # Preserve original extension spelling.
                suffix = path.suffix
                folder = [component({'year': day[:4], 'month': day[:7], **row}[key]) for key in rules.get('group_by', ['category'])]
                target = Path(*folder, stem + suffix)
                count = 1
                while str(target).casefold() in used:
                    count += 1
                    target = Path(*folder, f'{stem}__{count}{suffix}')
                used.add(str(target).casefold())
                row['target'] = target.as_posix()
                if digest(path) != checksum:
                    raise ValueError('扫描期间原文件发生变化，请重试')
                rows.append(row)
                hashes[checksum].append(str(path))
            except Exception as exc:
                errors.append({'path': str(path), 'reason': str(exc)})
    duplicates = {key: paths for key, paths in hashes.items() if len(paths) > 1}
    plan = dict(version=1, source=str(source), rules=rules, files=rows, duplicates=duplicates, errors=errors)
    save_json(output / 'plan.json', plan)
    save_json(output / 'duplicates.json', duplicates)
    save_json(output / 'errors.json', errors)
    report(output / 'preview.xlsx', rows, duplicates, errors)
    return plan


def apply(plan_path, output, confirmed=False):
    if not confirmed:
        raise ValueError('必须先检查预览，并明确确认后使用--confirm')
    plan = json.loads(Path(plan_path).read_text(encoding='utf-8'))
    if plan.get('version') != 1:
        raise ValueError('不支持的预览版本')
    source, output = Path(plan['source']).resolve(), Path(output).resolve()
    separate(source, output)
    # Validate all paths before creating output. Reject edited traversal plans.
    targets = set()
    for row in plan['files']:
        src, rel = Path(row['source']), Path(row['target'])
        if not src.is_absolute() or not src.is_relative_to(source) or src == source:
            raise ValueError('原件路径越界')
        if src.is_symlink() or src.resolve() != src or not src.resolve().is_relative_to(source):
            raise ValueError('原件路径包含符号链接或已越界')
        if rel.is_absolute() or not rel.parts or any(x in {'.', '..'} for x in rel.parts) or '\\' in str(rel):
            raise ValueError('目标路径越界')
        key = str(rel).casefold()
        if key in targets:
            raise ValueError('预览含重复目标路径')
        targets.add(key)
    output.mkdir(parents=True, exist_ok=False)
    files_root = output / 'files'
    files_root.mkdir()
    errors = list(plan['errors'])
    completed = []
    rows = []
    for original in plan['files']:
        row = dict(original)
        src, dst = Path(row['source']), files_root / row['target']
        created = False
        try:
            if src.resolve() != src:
                raise ValueError('原件路径已变化')
            dst.parent.mkdir(parents=True, exist_ok=True)
            checksum = hashlib.sha256()
            with regular_open(src) as read, dst.open('xb') as write:
                created = True
                for block in iter(lambda: read.read(1024 * 1024), b''):
                    checksum.update(block)
                    write.write(block)
            if checksum.hexdigest() != row['sha256'] or digest(dst) != row['sha256']:
                raise ValueError('原文件与预览不一致，或复制校验失败；请重新预览')
            os.utime(dst, (src.stat().st_atime, datetime.fromisoformat(row['date']).timestamp()))
            row['status'] = '已复制并校验'
            completed.append(dst)
        except Exception as exc:
            if created:
                dst.unlink(missing_ok=True)  # Only our new, failed output copy.
            row['status'] = '失败'
            row['warning'] = str(exc)
            errors.append({'path': str(src), 'reason': str(exc)})
        rows.append(row)
    save_json(output / 'errors.json', errors)
    save_json(output / 'duplicates.json', plan['duplicates'])
    report(output / 'inventory.xlsx', rows, plan['duplicates'], errors)
    with zipfile.ZipFile(output / 'organized.zip', 'x', compression=zipfile.ZIP_DEFLATED) as archive:
        for path in completed:
            archive.write(path, path.relative_to(output))
        for name in ('inventory.xlsx', 'duplicates.json', 'errors.json'):
            archive.write(output / name, name)
    return dict(copied=len(completed), failed=len(errors), output=str(output))


def main():
    parser = argparse.ArgumentParser(description='文件整理：先预览，确认后复制；永不移动或删除原件')
    sub = parser.add_subparsers(dest='command', required=True)
    prev = sub.add_parser('preview', help='只读扫描并生成Excel预览')
    prev.add_argument('source')
    prev.add_argument('--output', required=True)
    prev.add_argument('--rules')
    execute = sub.add_parser('apply', help='按已确认预览复制文件')
    execute.add_argument('plan')
    execute.add_argument('--output', required=True)
    execute.add_argument('--confirm', action='store_true')
    unpack = sub.add_parser('import-zip', help='安全导入作为输入容器的ZIP到新目录')
    unpack.add_argument('archive')
    unpack.add_argument('--output', required=True)
    ocr_pdf = sub.add_parser('ocr-pdf', help='确认后生成可搜索PDF副本，保留原件')
    ocr_pdf.add_argument('source')
    ocr_pdf.add_argument('--output', required=True)
    ocr_pdf.add_argument('--confirm', action='store_true')
    args = parser.parse_args()
    try:
        if args.command == 'preview':
            rules = json.loads(Path(args.rules).read_text(encoding='utf-8')) if args.rules else {}
            plan = preview(args.source, args.output, rules)
            result = dict(files=len(plan['files']), duplicate_groups=len(plan['duplicates']), failed=len(plan['errors']), content_warnings=sum(bool(r['warning']) for r in plan['files']), output=args.output)
        elif args.command == 'apply':
            result = apply(args.plan, args.output, args.confirm)
        elif args.command == 'ocr-pdf':
            from .ocr import searchable_pdf
            result = searchable_pdf(args.source, args.output, args.confirm)
        else:
            result = import_zip(args.archive, args.output)
        print(json.dumps(result, ensure_ascii=False))
        if result['failed']:
            sys.exit(2)
    except Exception as exc:
        print(f'操作失败：{exc}', file=sys.stderr)
        sys.exit(1)


def import_zip(archive_path, output):
    output = Path(output).resolve()
    archive_path = Path(archive_path).resolve()
    if archive_path.is_relative_to(output):
        raise ValueError('ZIP不能位于导入输出目录中')
    with zipfile.ZipFile(archive_path) as archive:
        entries = archive.infolist()
        if len(entries) > 10000 or sum(x.file_size for x in entries) > 2 * 1024**3:
            raise ValueError('ZIP条目数或总解压大小超过限制')
        seen = set()
        for item in entries:
            rel = Path(item.filename)
            mode = item.external_attr >> 16
            if (rel.is_absolute() or not rel.parts or '..' in rel.parts
                    or '\\' in item.filename or ':' in item.filename
                    or stat.S_ISLNK(mode) or item.flag_bits & 1
                    or item.file_size > 512 * 1024**2):
                raise ValueError(f'不安全或不支持的ZIP条目：{item.filename}')
            key = str(rel).casefold()
            if key in seen:
                raise ValueError('ZIP含重复路径')
            seen.add(key)
        output.mkdir(parents=True, exist_ok=False)
        errors = []
        copied = 0
        for item in entries:
            target = output / item.filename
            created = False
            try:
                if item.is_dir():
                    target.mkdir(parents=True, exist_ok=True)
                    continue
                target.parent.mkdir(parents=True, exist_ok=True)
                total = 0
                with archive.open(item) as read, target.open('xb') as write:
                    created = True
                    for block in iter(lambda: read.read(1024 * 1024), b''):
                        total += len(block)
                        if total > item.file_size or total > 512 * 1024**2:
                            raise ValueError('实际解压大小超过限制')
                        write.write(block)
                # ZIP timestamps do not carry timezone; interpret as Shanghai.
                stamp = datetime(*item.date_time, tzinfo=ZoneInfo('Asia/Shanghai')).timestamp()
                os.utime(target, (stamp, stamp))
                copied += 1
            except Exception as exc:
                if created:
                    target.unlink(missing_ok=True)
                errors.append({'path': item.filename, 'reason': str(exc)})
        # Keep diagnostics outside imported files so they aren't organized as input.
        if errors:
            import tempfile
            fd, log_path = tempfile.mkstemp(prefix='file-organizer-import-', suffix='.json', dir=output.parent)
            with os.fdopen(fd, 'w', encoding='utf-8') as stream:
                json.dump(errors, stream, ensure_ascii=False, indent=2)
        else:
            log_path = None
    return dict(imported=copied, failed=len(errors), error_log=log_path, output=str(output))
